"""Adversarial public-interface cases found during the October 2026 audit."""

import json
import subprocess
import sys

import pytest

from course_tutor.api import LumenCore
from course_tutor.errors import LumenError, ValidationError


@pytest.fixture
def core(tmp_path):
    with LumenCore(tmp_path / "workspace") as runtime:
        runtime.create_course("Synthetic audit", course_id="audit")
        yield runtime


def exercise(core, **kwargs):
    return core.create_exercise(
        "audit", "audit-family", "What is P(A|B)?",
        reference_solution="0.4", rubric={"max_points": 4}, **kwargs,
    )["exercise_version_id"]


def test_scoped_search_filters_before_limiting(core, tmp_path):
    wanted = core.create_topic("audit", "Wanted")
    sources = []
    for idx in range(3):
        path = tmp_path / f"lecture-{idx}.txt"
        path.write_text("### Theorem: Conditional probability\nContent: Bayes Bayes Bayes\n")
        sources.append(core.ingest_source(
            "audit", path, topic_slugs=[wanted["slug"]] if idx == 2 else [],
        )["data"])
    target = sources[-1]
    for query in ("Bayes", "הסתברות"):
        core.add_topic_alias("audit", wanted["id"], "הסתברות", "he")
        result = core.search("audit", query, source_id=target["source_id"], limit=1)
        assert [hit["unit_id"] for hit in result["results"]] == [target["units"][0]["id"]]
    result = core.search("audit", "Bayes", topic_id=wanted["id"], limit=1)
    assert result["results_count"] == 1
    assert core.search("audit", "Bayes", source_id="missing")["results"] == []


def test_ambiguous_unexpanded_title_is_rejected(core, tmp_path):
    for idx in range(2):
        path = tmp_path / f"version-{idx}.txt"
        path.write_text(f"### Theorem: Bayes\nContent: version {idx}\n")
        core.ingest_source("audit", path)
    with pytest.raises(ValidationError, match="Ambiguous"):
        core.read_units("audit", ["Bayes"])


def test_missing_original_is_reported_as_missing(core, tmp_path):
    path = tmp_path / "lecture.txt"
    path.write_text("Synthetic lecture")
    source = core.ingest_source("audit", path)["data"]
    (core.workspace_root / source["storage_relpath"]).unlink()
    region = core.read_source_region(source["source_version_id"])
    assert region["file_exists"] is False
    assert region["hash_verified"] is False
    assert region["absolute_path"] is None


def test_cross_course_evidence_rolls_back(core):
    core.create_course("Other", course_id="other")
    foreign = core.create_topic("other", "Foreign")
    attempt = core.start_attempt("audit", exercise(core))["data"]
    core.record_attempt_event(attempt["id"], 1, "work_submission", content="0.4")
    with pytest.raises(LumenError):
        core.commit_evidence(attempt["id"], [{
            "topic_id": foreign["id"], "skill_id": "conditional", "correct": True,
        }], request_id="cross-course")
    assert core.evidence_repo.list_evidence("audit") == []
    assert core.con.execute("SELECT COUNT(*) FROM operation_receipts WHERE request_id = 'cross-course'").fetchone()[0] == 0


@pytest.mark.parametrize("operation", ["parent", "alias", "link", "exercise"])
def test_cross_course_topic_relationships_rejected(core, operation):
    core.create_course("Other", course_id="other")
    foreign = core.create_topic("other", "Foreign")
    own = core.create_topic("audit", "Own")
    with pytest.raises(LumenError):
        if operation == "parent":
            core.create_topic("audit", "Child", parent_topic_id=foreign["id"])
        elif operation == "alias":
            core.add_topic_alias("audit", foreign["id"], "foreign")
        elif operation == "link":
            core.add_topic_link("audit", own["id"], foreign["id"])
        else:
            exercise(core, topic_mappings=[foreign["id"]])


def test_diagnostic_respects_rubric_maximum(core):
    version_id = exercise(core)
    core.validate_exercise_version(version_id, checks={"reviewed": True}, expected_revision=1)
    assessment = core.create_assessment("audit", "Diagnostic", {}, [version_id])["data"]
    item = core.assessment_progress(assessment["assessment_id"])["items"][0]
    assert item["max_points"] == 4


def test_create_exercise_failure_is_atomic(core, monkeypatch):
    def fail(*args, **kwargs):
        raise ValidationError("simulated version write failure")
    monkeypatch.setattr(core.exercise_repo, "create_exercise_version", fail)
    with pytest.raises(ValidationError):
        exercise(core)
    assert core.exercise_repo.list_exercises("audit") == []


def test_non_idempotent_serialization_failure_rolls_back(core, monkeypatch):
    original = core.course_repo.create_course
    def mutate_response(course):
        original(course)
        course.notation_settings = {"bad": float("nan")}
    monkeypatch.setattr(core.course_repo, "create_course", mutate_response)
    with pytest.raises(ValidationError):
        core.create_course("Must roll back", course_id="bad-response")
    assert core.course_repo.get_course("bad-response") is None


@pytest.mark.parametrize("destination", ["workspace", "originals", "ancestor", "database"])
def test_backup_rejects_destructive_destination(core, destination):
    paths = {
        "workspace": core.workspace_root,
        "originals": core.workspace_root / "originals",
        "ancestor": core.workspace_root.parent,
        "database": core.db_path,
    }
    with pytest.raises(LumenError):
        core.backup(paths[destination])
    assert core.db_path.is_file()
    assert core.course_brief("audit")["course"]["id"] == "audit"


def test_ingest_retry_conflict_cleans_staging(core, tmp_path):
    path = tmp_path / "lecture.txt"
    path.write_text("First synthetic source")
    core.ingest_source("audit", path, request_id="ingest")
    path.write_text("Different synthetic source")
    with pytest.raises(LumenError):
        core.ingest_source("audit", path, request_id="ingest")
    assert list((core.workspace_root / "staging").iterdir()) == []


def test_cli_preserves_assessment_owner(core):
    args = [sys.executable, "-m", "course_tutor.cli", "create-assessment", "-w", str(core.workspace_root),
            "-c", "audit", "-j", json.dumps({"title": "Owned", "learner_id": "alice", "blueprint": {}})]
    result = subprocess.run(args, capture_output=True, text=True, check=True)
    assessment_id = json.loads(result.stdout)["data"]["assessment_id"]
    assert core.assessment_repo.get_assessment(assessment_id).learner_id == "alice"


def test_hebrew_identity_agrees_across_resume_and_reporting(core):
    version_id = core.create_exercise(
        "audit", "hebrew", "חשבו הסתברות מותנית", reference_solution="הפתרון: 1/2",
        rubric={"max_points": 4, "criterion": "חישוב נכון"},
    )["exercise_version_id"]
    attempt = core.start_attempt("audit", version_id)["data"]
    assert core.course_brief("audit")["open_attempt"]["attempt_id"] == attempt["id"]
    core.resume_attempt(attempt["id"], 1)


def test_images_remain_originals_with_explicit_extraction_uncertainty(core, tmp_path):
    # A real 1x1 PNG; no OCR path exists in this application.
    import base64
    path = tmp_path / "student-work.png"
    path.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/l9sAAAAASUVORK5CYII="))
    source = core.ingest_source("audit", path)["data"]
    assert source["extraction_status"] == "uncertain"
    assert source["limitations"]
    assert source["units"] == []
    assert core.read_source_region(source["source_version_id"])["hash_verified"] is True


def test_reused_corrupt_original_is_not_acknowledged(core, tmp_path):
    path = tmp_path / "lecture.txt"
    path.write_text("Original synthetic bytes")
    source = core.ingest_source("audit", path)["data"]
    (core.workspace_root / source["storage_relpath"]).write_text("Corrupted")
    with pytest.raises(LumenError):
        core.ingest_source("audit", path)


def test_publish_false_keeps_enriched_units_pending(core, tmp_path):
    path = tmp_path / "lecture.txt"
    path.write_text("Synthetic source")
    source = core.ingest_source("audit", path)["data"]
    enriched = core.enrich_source_version(source["source_version_id"], [
        {"title": "Unreviewed", "content": "Generated interpretation"},
    ], publish=False)
    assert enriched["units"][0]["status"] == "pending_review"


def test_prior_solution_exposure_survives_a_new_attempt(core):
    topic = core.create_topic("audit", "Conditional")
    version_id = exercise(core)
    old = core.start_attempt("audit", version_id)["data"]
    core.record_attempt_event(old["id"], 1, "solution_revealed", content="0.4")
    new = core.start_attempt("audit", version_id)["data"]
    core.record_attempt_event(new["id"], 1, "work_submission", content="0.4")
    result = core.commit_evidence(new["id"], [{"topic_id": topic["id"], "skill_id": "conditional", "correct": True}])["data"]
    assert result["evidence_created"][0]["demonstrated_level"] == "assisted_success"
    with pytest.raises(ValidationError):
        core.correct_evidence(result["evidence_created"][0]["id"], "Must retain prior exposure", {"demonstrated_level": "independent_mastery"})


def test_false_completion_flag_cannot_complete_attempt(core):
    attempt = core.start_attempt("audit", exercise(core))["data"]
    with pytest.raises(ValidationError):
        core.commit_evidence(attempt["id"], [], complete_attempt="false")
    assert core.attempt_repo.get_attempt(attempt["id"]).status == "open"


def test_learner_settings_persist_through_public_api(core):
    saved = core.update_learner_settings("audit", learner_id="alice", chosen_language="he",
                                       teaching_preference="no hints unless requested", request_id="settings")
    with LumenCore(core.workspace_root) as reopened:
        brief = reopened.course_brief("audit", learner_id="alice")
        assert brief["learner_settings"]["chosen_language"] == "he"
        assert brief["learner_settings"]["teaching_preference"] == "no hints unless requested"
        assert reopened.update_learner_settings("audit", learner_id="alice", chosen_language="he",
                                               teaching_preference="no hints unless requested", request_id="settings")["idempotent_replay"]
    assert saved["receipt"]


def test_slide_tables_and_visual_uncertainty_are_preserved(core, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from io import BytesIO
    import base64
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "P(A)", "P(B)"
    table.cell(1, 0).text, table.cell(1, 1).text = "0.2", "0.5"
    slide.shapes.add_picture(BytesIO(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/l9sAAAAASUVORK5CYII=")), Inches(5), Inches(1))
    path = tmp_path / "table.pptx"
    deck.save(path)
    source = core.ingest_source("audit", path)["data"]
    assert source["extraction_status"] == "uncertain"
    sections = core.get_extracted_sections(source["source_version_id"])
    assert "P(A)\tP(B)" in sections[0]["text"]
    assert "0.2\t0.5" in sections[0]["text"]
    assert sections[0]["metadata"]["has_images"]


def test_capabilities_does_not_overwrite_existing_file(tmp_path):
    from course_tutor.capabilities import check_capabilities
    sentinel = tmp_path / ".lumen_cap_test"
    sentinel.write_text("Existing user data")
    assert check_capabilities(tmp_path)["workspace"]["writable"]
    assert sentinel.read_text() == "Existing user data"


def test_imported_ocr_library_does_not_claim_an_unimplemented_pipeline(monkeypatch):
    import sys
    from course_tutor.capabilities import check_capabilities
    monkeypatch.setitem(sys.modules, "pytesseract", object())
    assert check_capabilities()["extraction_capabilities"]["ocr_scans"] is False


def test_duplicate_content_reuse_is_scoped_to_logical_source(core, tmp_path):
    path = tmp_path / "lecture.txt"
    path.write_text("Identical bytes reused by multiple logical sources")
    first = core.ingest_source("audit", path, logical_name="first")["data"]
    core.ingest_source("audit", path, logical_name="second")
    replay = core.ingest_source("audit", path, logical_name="first")["data"]
    assert replay["source_version_id"] == first["source_version_id"]
    assert replay["version_num"] == 1


def test_repeated_family_does_not_inflate_fresh_independent_evidence(core):
    topic = core.create_topic("audit", "Conditional")
    version_id = exercise(core)
    for _ in range(2):
        attempt = core.start_attempt("audit", version_id)["data"]
        core.record_attempt_event(attempt["id"], 1, "work_submission", content="0.4")
        core.commit_evidence(attempt["id"], [{"topic_id": topic["id"], "skill_id": "conditional", "correct": True}])
    state = core.evidence_repo.get_topic_state("audit", topic["id"])
    assert len(core.evidence_repo.list_evidence("audit")) == 2  # observations remain auditable
    assert state.independent_success_count == 1
    assert state.summary_data["independent_observation_count"] == 2


@pytest.mark.parametrize("operation", ["create_topic", "add_topic_alias", "add_topic_link", "create_exercise",
                                       "enrich_source_version", "validate_exercise_version", "update_exercise_version"])
def test_authoring_mutations_replay_without_duplicate_writes(core, tmp_path, operation):
    topic = core.create_topic("audit", "Conditional")
    other = core.create_topic("audit", "Other")
    version_id = exercise(core)
    path = tmp_path / "source.txt"
    path.write_text("Synthetic lecture")
    source = core.ingest_source("audit", path)["data"]
    parameters = {
        "create_topic": {"course_id": "audit", "name": "New topic"},
        "add_topic_alias": {"course_id": "audit", "topic_id": topic["id"], "alias": "תנאי"},
        "add_topic_link": {"course_id": "audit", "from_id": topic["id"], "to_id": other["id"]},
        "create_exercise": {"course_id": "audit", "exercise_family": "new", "question_text": "Generated question"},
        "enrich_source_version": {"source_version_id": source["source_version_id"], "structured_units": [{"title": "Reviewed", "content": "Synthetic"}]},
        "validate_exercise_version": {"version_id": version_id, "expected_revision": 1, "checks": {"reviewed": True}},
        "update_exercise_version": {"version_id": version_id, "expected_revision": 1, "question_text": "Revised synthetic question"},
    }[operation]
    call = getattr(core, operation)
    first = call(**parameters, request_id="authoring")
    replay = call(**parameters, request_id="authoring")
    assert replay["idempotent_replay"]
    assert replay["data"] == first["data"]
    assert replay["receipt"] == first["receipt"]


def test_legacy_projection_rebuilds_without_erasing_observations(core):
    topic = core.create_topic("audit", "Conditional")
    version_id = exercise(core)
    for _ in range(2):
        attempt = core.start_attempt("audit", version_id)["data"]
        core.record_attempt_event(attempt["id"], 1, "work_submission", content="0.4")
        core.commit_evidence(attempt["id"], [{"topic_id": topic["id"], "skill_id": "conditional", "correct": True}])
    core.con.execute("UPDATE topic_state SET policy_version=1, independent_success_count=99")
    rebuilt = core.evidence_repo.get_topic_state("audit", topic["id"])
    assert rebuilt.policy_version == 2
    assert rebuilt.independent_success_count == 1
    assert len(core.evidence_repo.list_evidence("audit")) == 2


@pytest.mark.parametrize("path", ["C:\\originals\\lecture.txt", "C:lecture.txt", "\\\\server\\share\\lecture.txt"])
def test_restore_rejects_windows_paths_on_every_platform(tmp_path, path):
    from course_tutor.backup.restore import _safe_resolve_staged_path
    with pytest.raises(LumenError):
        _safe_resolve_staged_path(tmp_path, path)


def test_contradictory_correctness_and_mastery_cannot_create_positive_evidence(core):
    topic = core.create_topic("audit", "Conditional")
    attempt = core.start_attempt("audit", exercise(core))["data"]
    core.record_attempt_event(attempt["id"], 1, "work_submission", content="Incorrect answer")
    with pytest.raises(ValidationError):
        core.commit_evidence(attempt["id"], [{"topic_id": topic["id"], "skill_id": "conditional",
                                              "correct": False, "demonstrated_level": "independent_mastery"}])
    assert core.evidence_repo.list_evidence("audit") == []


def test_course_brief_preserves_visible_partial_progress(core):
    """A real submitted solution had correct intermediate work and a wrong choice."""
    topic = core.create_topic("audit", "Correct intermediate method")
    attempt = core.start_attempt("audit", exercise(core))["data"]
    core.record_attempt_event(attempt["id"], 1, "work_submission", content="Correct method; final choice incorrect")
    core.commit_evidence(attempt["id"], [{"topic_id": topic["id"], "skill_id": "method",
                                         "correct": True, "partial": True,
                                         "demonstrated_level": "partial_progress"}])
    summary = next(t for t in core.course_brief("audit")["coverage"]["assessed_topics"] if t["name"] == topic["name"])
    assert summary["partial_progress"] == 1
    assert summary["independent"] == 0
    assert summary["errors"] == 0
