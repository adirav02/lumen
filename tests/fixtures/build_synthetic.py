"""Build tiny generated binary fixtures offline; all content is synthetic."""
from pathlib import Path
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from pptx import Presentation
from PIL import Image, ImageDraw

ROOT = Path(__file__).parent


def pdf(path, pages):
    writer = PdfWriter()
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
            DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        escaped = [line.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)') for line in lines]
        stream = DecodedStreamObject()
        stream.set_data(('BT /F1 12 Tf 40 750 Td 18 TL\n'+'\n'.join('('+line+') Tj T*' for line in escaped)+'\nET').encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(path)


def slides(path):
    deck = Presentation()
    for title, text in [('SYNTHETIC: Independence', 'אי תלות: P(A intersection B)=P(A)P(B).'),
                        ('SYNTHETIC: Variables', 'A random variable maps outcomes to numbers.')]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = text
    deck.save(path)


def build():
    audit = ROOT/'audit_course/materials'
    prob = ROOT/'synthetic_probability/materials'
    pdf(audit/'conditional_probability.pdf', [
        ['SYNTHETIC: Conditional probability requires P(B)>0.', 'P(A|B)=P(A intersection B)/P(B).'],
        ['SYNTHETIC: Bayes. P(D)=0.02; sensitivity=0.90; false positives=0.10.', 'Posterior = 9/58.']])
    pdf(prob/'sample_lecture.pdf', [['SYNTHETIC: Bayes Rule.', 'Assume P(B) > 0.', 'P(A|B)=P(A intersection B)/P(B).']])
    slides(audit/'independence_slides.pptx')
    slides(prob/'sample_slides.pptx')
    for name, middle in [('graph.png', '1'), ('graph-unknown.png', '?')]:
        image = Image.new('RGB', (320, 200), 'white')
        draw = ImageDraw.Draw(image)
        points = {'s': (25,100), 'a': (125,30), 'b': (125,170), 't': (285,100)}
        for a,b,weight in [('s','a','2'),('s','b','5'),('a','b',middle),('a','t','7'),('b','t','2')]:
            p,q=points[a],points[b]; draw.line([p,q],fill='black',width=2)
            draw.text(((p[0]+q[0])/2,(p[1]+q[1])/2),weight,fill='blue')
        for label,point in points.items():
            draw.text(point,label,fill='black')
        image.save(ROOT/'educational_evaluation'/name)
