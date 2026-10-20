"""Rebuild generated text PDF and PPTX fixtures; never invoke application code."""
from pathlib import Path

from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from pptx import Presentation


ROOT = Path(__file__).parent / "materials"


def build_pdf():
    writer = PdfWriter()
    pages = [
        ["SYNTHETIC GENERATED AUDIT MATERIAL - page 1",
         "Conditional probability requires P(B)>0.",
         "P(A|B)=P(A intersection B)/P(B).",
         "If P(B)=0, this elementary conditional is undefined.",
         "On two fair independent coins: P(A)=P(B)=1/2.",
         "For first-heads A and second-heads B: P(A intersection B)=1/4."],
        ["SYNTHETIC GENERATED AUDIT MATERIAL - page 2",
         "Bayes v1: P(D)=0.02; P(+|D)=0.90; P(+|not-D)=0.10.",
         "D and not-D partition the population.",
         "P(+)=0.018+0.098=0.116>0.",
         "P(D|+)=0.018/0.116=9/58.",
         "Version 2 will change false-positive rate to 0.05; posterior 18/67."],
    ]
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    font_ref = writer._add_object(font)
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})})
        def escape(line):
            return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content = "BT /F1 12 Tf 40 750 Td 18 TL\n" + "\n".join(f"({escape(line)}) Tj T*" for line in lines) + "\nET"
        stream = DecodedStreamObject()
        stream.set_data(content.encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = ROOT / "conditional_probability.pdf"
    with output.open("wb") as handle:
        writer.write(handle)
    assert len(PdfReader(output).pages) == 2
    assert "requires P(B)>0" in PdfReader(output).pages[0].extract_text()


def build_slides():
    deck = Presentation()
    for title, text in [
        ("SYNTHETIC AUDIT: Independence", "Independent iff P(A intersection B)=P(A)P(B).\nWith P(B)>0: P(A|B)=P(A).\nאי תלות = independence; הסתברות מותנית = conditional probability."),
        ("SYNTHETIC AUDIT: Counterexample", "One fair die: A={1}, B={2}.\nP(A intersection B)=0; P(A)P(B)=1/36.\nDisjoint positive-probability events are dependent.\nEmpty event is independent of every event by the product definition."),
    ]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = text
    output = ROOT / "independence_slides.pptx"
    deck.save(output)
    assert len(Presentation(output).slides) == 2


if __name__ == "__main__":
    build_pdf()
    build_slides()
